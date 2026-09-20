"""
Single-label Excel workbook builder.

Architectural responsibility: serialize one verification response to .xlsx bytes.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment

from app.models.verification import SingleReviewVerificationResponse
from app.reporting.sanitize import sanitize_spreadsheet_value
from app.reporting.styles import (
    BODY_FONT,
    DISCLAIMER_FONT,
    PROTOTYPE_DISCLAIMER,
    SECTION_FONT,
    style_header_row,
    style_title_row,
    write_kv,
)

FIELD_ORDER = (
    ("brand_name", "Brand Name"),
    ("class_type", "Class / Type"),
    ("alcohol_content", "Alcohol Content"),
    ("net_contents", "Net Contents"),
    ("government_warning", "Government Warning"),
)


def build_single_review_xlsx(result: SingleReviewVerificationResponse) -> bytes:
    """Return Office Open XML bytes for a Single Label review."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Review Summary"

    style_title_row(
        ws,
        1,
        2,
        "LabelVerify — Single Label Review",
        "Prototype · AI-Assisted Review · Decision support only",
    )

    row = 4
    ws.cell(row=row, column=1, value="Review metadata").font = SECTION_FONT
    row += 1
    row = write_kv(ws, row, "Review ID", result.verification_id)
    row = write_kv(ws, row, "Generated (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S"))
    row = write_kv(ws, row, "Filename", result.filename)
    row = write_kv(ws, row, "Verification mode", result.verification_mode)
    row = write_kv(ws, row, "Overall finding", result.verification.overall_status.value)
    row = write_kv(ws, row, "Processing time (ms)", f"{result.processing_time_ms:.1f}")
    row = write_kv(ws, row, "AI-assisted evidence", "Yes" if result.verification.ai_used else "No")
    row = write_kv(ws, row, "Passed", result.verification.pass_count)
    row = write_kv(ws, row, "Needs review", result.verification.review_count)
    row = write_kv(ws, row, "Failed", result.verification.fail_count)

    row += 1
    ws.cell(row=row, column=1, value="Extracted label information").font = SECTION_FONT
    row += 1
    fields = result.extraction.selected_fields
    for key, label in FIELD_ORDER:
        field = fields.get(key)
        if field is None:
            value = "Not extracted"
        elif key == "government_warning":
            if field.status.value == "FOUND":
                value = "Detected"
                if field.raw_text:
                    value = f"Detected — {field.raw_text[:200]}"
            else:
                value = f"{field.status.value}: {field.raw_text or field.explanation or ''}"
        else:
            value = field.raw_text or field.normalized_value or field.status.value
        row = write_kv(ws, row, label, value)

    row += 2
    ws.cell(row=row, column=1, value="Verification checks").font = SECTION_FONT
    row += 1
    headers = [
        "Check",
        "Status",
        "Detected Evidence",
        "Explanation",
        "Reason Code",
        "AI Assisted",
    ]
    style_header_row(ws, row, headers)
    row += 1
    for check in result.verification.checks:
        ai_flag = bool((check.technical_details or {}).get("ai_assisted_evidence"))
        values = [
            check.check_name,
            check.status.value,
            check.detected_label_value or "",
            check.explanation,
            check.reason_code or "",
            "Yes" if ai_flag else "No",
        ]
        for col, raw in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=sanitize_spreadsheet_value(raw))
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row += 1
    disc = ws.cell(row=row, column=1, value=PROTOTYPE_DISCLAIMER)
    disc.font = DISCLAIMER_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 36
    ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 14

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
