"""
Batch Excel workbook builder.

Architectural responsibility: serialize a batch job record to multi-sheet .xlsx bytes.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment

from app.batch.models import BatchItemProcessingState, BatchSummaryCounts
from app.reporting.sanitize import sanitize_spreadsheet_value
from app.reporting.styles import (
    BODY_FONT,
    DISCLAIMER_FONT,
    PROTOTYPE_DISCLAIMER,
    SECTION_FONT,
    style_header_row,
    style_title_row,
    write_kv,
)

if TYPE_CHECKING:
    from app.batch.store import BatchJobRecord

RESULT_HEADERS = [
    "Filename",
    "Overall Status",
    "Brand",
    "Class / Type",
    "Alcohol Content",
    "Net Contents",
    "Brand Check",
    "Class / Type Check",
    "Alcohol Check",
    "Net Contents Check",
    "Warning Check",
    "AI Assisted",
    "Review Reason Codes",
    "Processing Time (ms)",
    "Error",
]


def build_batch_xlsx(job: BatchJobRecord) -> bytes:
    """Return Office Open XML bytes for a Batch Review."""
    wb = Workbook()

    summary = wb.active
    summary.title = "Batch Summary"
    style_title_row(
        summary,
        1,
        2,
        "LabelVerify — Batch Review",
        "Prototype · AI-Assisted Review · Decision support only",
    )
    row = 4
    summary.cell(row=row, column=1, value="Batch metadata").font = SECTION_FONT
    row += 1
    s = _summary_counts(job)
    row = write_kv(summary, row, "Batch ID", job.batch_id)
    row = write_kv(summary, row, "Generated (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S"))
    row = write_kv(summary, row, "State", job.state.value)
    row = write_kv(summary, row, "Total", s.total)
    row = write_kv(summary, row, "PASS", s.pass_count)
    row = write_kv(summary, row, "REVIEW", s.review_count)
    row = write_kv(summary, row, "FAIL", s.fail_count)
    row = write_kv(summary, row, "ERROR", s.error_count)
    row = write_kv(summary, row, "AI-assisted count", s.ai_assisted_count)
    duration = None
    if job.started_at is not None and job.finished_at is not None:
        duration = max(0.0, (job.finished_at - job.started_at) * 1000)
    row = write_kv(
        summary,
        row,
        "Processing duration (ms)",
        f"{duration:.1f}" if duration is not None else "",
    )
    row += 1
    disc = summary.cell(row=row, column=1, value=PROTOTYPE_DISCLAIMER)
    disc.font = DISCLAIMER_FONT
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 48

    results = wb.create_sheet("Review Results")
    style_title_row(
        results,
        1,
        len(RESULT_HEADERS),
        "LabelVerify — Batch Review Results",
        "One row per label · Prototype decision support",
    )
    header_row = 4
    style_header_row(results, header_row, RESULT_HEADERS)
    data_row = header_row + 1
    for item in job.items:
        checks = _check_statuses(item)
        values = [
            item.filename,
            item.overall_status or item.processing_state.value,
            item.row.brand_name,
            item.row.class_type,
            item.row.alcohol_content,
            item.row.net_contents,
            checks.get("Brand Name", ""),
            checks.get("Class / Type", ""),
            checks.get("Alcohol Content / ABV", ""),
            checks.get("Net Contents", ""),
            checks.get("Government Health Warning", ""),
            "Yes" if item.ai_assisted else "No",
            ";".join(item.review_reason_codes),
            f"{item.processing_time_ms:.1f}" if item.processing_time_ms is not None else "",
            item.error_message or "",
        ]
        for col, raw in enumerate(values, start=1):
            cell = results.cell(
                row=data_row,
                column=col,
                value=sanitize_spreadsheet_value(raw),
            )
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        data_row += 1

    # Extend autofilter to data
    from openpyxl.utils import get_column_letter

    end = get_column_letter(len(RESULT_HEADERS))
    results.auto_filter.ref = f"A{header_row}:{end}{max(data_row - 1, header_row)}"

    widths = [22, 14, 22, 28, 16, 14, 12, 14, 12, 12, 12, 12, 28, 14, 36]
    for idx, width in enumerate(widths, start=1):
        results.column_dimensions[get_column_letter(idx)].width = width

    disc_row = data_row + 1
    dcell = results.cell(row=disc_row, column=1, value=PROTOTYPE_DISCLAIMER)
    dcell.font = DISCLAIMER_FONT

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _summary_counts(job: BatchJobRecord) -> BatchSummaryCounts:
    counts = BatchSummaryCounts(total=len(job.items))
    for item in job.items:
        if item.processing_state == BatchItemProcessingState.QUEUED:
            counts.queued_count += 1
        elif item.processing_state == BatchItemProcessingState.PROCESSING:
            counts.processing_count += 1
        elif item.processing_state == BatchItemProcessingState.ERROR:
            counts.error_count += 1
        elif item.processing_state == BatchItemProcessingState.COMPLETED:
            counts.completed += 1
            if item.overall_status == "PASS":
                counts.pass_count += 1
            elif item.overall_status == "REVIEW":
                counts.review_count += 1
            elif item.overall_status == "FAIL":
                counts.fail_count += 1
        if item.ai_assisted:
            counts.ai_assisted_count += 1
    counts.ai_budget_exhausted = job.ai_budget_exhausted
    counts.ai_calls_used = job.ai_calls_used
    return counts


def _check_statuses(item: object) -> dict[str, str]:
    result = getattr(item, "result", None)
    if result is None:
        return {}
    verification = result.verification
    return {c.check_name: c.status.value for c in verification.checks}
