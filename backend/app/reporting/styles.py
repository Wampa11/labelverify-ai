"""
Shared openpyxl styling for LabelVerify prototype reports.

Architectural responsibility: restrained institutional formatting — not official Treasury branding.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# Institutional navy-adjacent fills (hex); not official seals/insignia.
NAVY_FILL = PatternFill("solid", fgColor="1B3A5C")
HEADER_FILL = PatternFill("solid", fgColor="E8EEF4")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="FFFFFF")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1B3A5C")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="334155")
BODY_FONT = Font(name="Calibri", size=10, color="1E293B")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="1B3A5C")
DISCLAIMER_FONT = Font(name="Calibri", size=9, italic=True, color="64748B")
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

PROTOTYPE_DISCLAIMER = (
    "Prototype decision-support output. Final regulatory determinations remain with "
    "authorized personnel. Not an official U.S. Department of the Treasury report."
)


def style_title_row(ws: Worksheet, row: int, cols: int, title: str, subtitle: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = TITLE_FONT
    cell.fill = NAVY_FILL
    cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[row].height = 24
    sub = row + 1
    ws.merge_cells(start_row=sub, start_column=1, end_row=sub, end_column=cols)
    sc = ws.cell(row=sub, column=1, value=subtitle)
    sc.font = SUBTITLE_FONT
    sc.fill = NAVY_FILL
    ws.row_dimensions[sub].height = 18


def style_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.auto_filter.ref = f"A{row}:{_col_letter(len(headers))}{row}"
    ws.freeze_panes = f"A{row + 1}"


def write_kv(ws: Worksheet, row: int, label: str, value: object) -> int:
    from app.reporting.sanitize import sanitize_spreadsheet_value

    lc = ws.cell(row=row, column=1, value=label)
    lc.font = LABEL_FONT
    vc = ws.cell(row=row, column=2, value=sanitize_spreadsheet_value(value))
    vc.font = BODY_FONT
    vc.alignment = Alignment(wrap_text=True)
    return row + 1


def _col_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)
